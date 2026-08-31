import pandas as pd
from openpyxl import Workbook, load_workbook

from backend.engine.core.loader import (
    _pick_pn_column,
    _prepare_france_price_file,
    _prepare_sys_price_file,
    _read_excel_any,
)


def _write_report_price(path):
    wb = Workbook()
    nav = wb.active
    nav.title = "navigation"
    nav.append(["", ""])

    products = wb.create_sheet("products")
    products.append(["", "Back to Navigation", None])
    products.append(["Part No.", "Series", "FOB C(EUR)"])
    products.append(["1.0.01", "Cabling", "12.30"])
    products.append([None, None, None])
    wb.save(path)


def test_prepare_france_price_consumes_report_price_products_sheet(tmp_path):
    data_dir = tmp_path / "data"
    data_dir.mkdir()

    old_france = data_dir / "FrancePrice.xlsx"
    old_wb = Workbook()
    old_wb.active.append(["Part No.", "Series"])
    old_wb.active.append(["old", "old"])
    old_wb.save(old_france)

    report = data_dir / "reportPrice_1782735392298.xlsx"
    _write_report_price(report)

    france_path = _prepare_france_price_file(data_dir)

    assert france_path == old_france
    assert not report.exists()

    wb = load_workbook(france_path, read_only=True, data_only=True)
    assert wb.sheetnames == ["products"]
    ws = wb["products"]
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0] == ("Part No.", "Series", "FOB C(EUR)")
    assert rows[1] == ("1.0.01", "Cabling", 12.3)


def test_prepare_sys_price_consumes_pricelist_export(tmp_path):
    data_dir = tmp_path / "data"
    data_dir.mkdir()

    old_sys = data_dir / "SysPrice.xlsx"
    old_wb = Workbook()
    old_wb.active.append(["Part Num", "Internal Model"])
    old_wb.active.append(["old", "old"])
    old_wb.save(old_sys)

    price_list = data_dir / "(20260629122325) PriceList.xls"
    wb = Workbook()
    ws = wb.active
    ws.append(["Part Num", "Internal Model", "Min Price", "Area Price", "Sales Type"])
    ws.append(["1.0.02", "DHI-TEST", 10, 20, "Distribution"])
    wb.save(price_list)

    sys_path = _prepare_sys_price_file(data_dir)

    assert sys_path == old_sys
    assert not price_list.exists()

    wb = load_workbook(sys_path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0] == ("Part Num", "Internal Model", "Min Price", "Area Price", "Sales Type")
    assert rows[1] == ("1.0.02", "DHI-TEST", 10, 20, "Distribution")


def test_read_excel_discovers_pn_table_after_navigation_sheet(tmp_path):
    report = tmp_path / "renamed-upload.xlsx"
    _write_report_price(report)

    df = _read_excel_any(report)

    assert list(df.columns) == ["Part No.", "Series", "FOB C(EUR)"]
    assert df.to_dict("records") == [
        {"Part No.": "1.0.01", "Series": "Cabling", "FOB C(EUR)": 12.3}
    ]


def test_pick_pn_column_tolerates_export_header_formatting():
    df = pd.DataFrame(columns=["Series", "Part\nNo.", "Price"])

    assert _pick_pn_column(df) == "Part\nNo."
